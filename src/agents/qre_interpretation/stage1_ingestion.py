"""Stage 1 — ingestion. DOCX to a document object in true body order.

Iterates the body XML so headings, paragraphs and tables stay interleaved;
`doc.paragraphs` and `doc.tables` are separate sequences and lose that ordering.

No LLM. Literal transcription only, with one derivation: where a document
uses no heading styles at all, its headings are recognised by how they are set
(see `_inferred_headings`). Without that such a document has no sections, and
every later stage finds nothing in it.
"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from .models import BlockKind, Paragraph, Stage1Document, Table

_HEADING_LEVEL = re.compile(r"^heading\s+(\d+)$", re.IGNORECASE)


class IngestionError(Exception):
    """File missing, wrong extension, or unreadable. Never a partial document."""


def _heading_level(style: str) -> int | None:
    match = _HEADING_LEVEL.match(style.strip())
    return int(match.group(1)) if match else None


#: A heading is a line, not a paragraph of prose.
_MAX_HEADING_CHARS = 120


def _set_as(paragraph: DocxParagraph) -> tuple[bool, float | None]:
    """(every run bold, the one font size used) for a paragraph's visible text.

    The size is None where runs disagree or none states one outright.
    """
    runs = [run for run in paragraph.runs if run.text.strip()]
    sizes = {run.font.size.pt if run.font.size else None for run in runs}
    size = sizes.pop() if len(sizes) == 1 else None
    return bool(runs) and all(run.bold for run in runs), size


def _inferred_headings(items: list) -> set[int]:
    """Body positions of headings in a document that styles none.

    Some QREs set their headings by hand: a line wholly in bold, larger than
    the body text, in the default paragraph style. That is typography, not
    wording, so it is read the same way whatever the headings say. A size used
    by a single such line is a title or a cover line rather than a level of an
    outline, and stays out — the cover has to remain front matter.

    Empty when the document does use heading styles: then the styles are the
    author's own statement of structure and nothing is inferred beside them.
    """
    paragraphs = [(i, p) for i, p in enumerate(items) if isinstance(p, DocxParagraph)]
    if any(
        _heading_level(p.style.name) is not None
        for _, p in paragraphs
        if p.style is not None
    ):
        return set()

    body_sizes: Counter = Counter()
    candidates: list[tuple[int, float]] = []
    for index, paragraph in paragraphs:
        text = paragraph.text.strip()
        bold, size = _set_as(paragraph)
        if size is None:
            continue
        body_sizes[size] += len(text)
        if bold and len(text) <= _MAX_HEADING_CHARS:
            candidates.append((index, size))
    if not body_sizes:
        return set()

    body = body_sizes.most_common(1)[0][0]
    repeated = Counter(size for _, size in candidates if size > body)
    return {index for index, size in candidates if repeated[size] > 1}


def _text(raw: str) -> str:
    """Text as written, with every line break as a newline.

    A soft line break can arrive as U+2028 or U+2029 instead of "\n". Left as
    they are, a cell holding three instructions on three lines reads downstream
    as one line.
    """
    return raw.replace("\u2028", "\n").replace("\u2029", "\n")


def _iter_body(document) -> list[DocxParagraph | DocxTable]:
    items: list[DocxParagraph | DocxTable] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            items.append(DocxParagraph(child, document))
        elif child.tag == qn("w:tbl"):
            items.append(DocxTable(child, document))
    return items


def run(path: str | Path) -> Stage1Document:
    path = Path(path)
    if not path.exists():
        raise IngestionError(f"File not found: {path}")
    if path.suffix.lower() != ".docx":
        raise IngestionError(f"Expected .docx, got '{path.suffix}': {path}")

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"Could not read '{path}': {exc}") from exc

    items = _iter_body(document)
    inferred = _inferred_headings(items)

    blocks: list[Paragraph | Table] = []
    for order, item in enumerate(items):
        if isinstance(item, DocxParagraph):
            style = item.style.name if item.style is not None else "Normal"
            blocks.append(
                Paragraph(
                    kind=BlockKind.PARAGRAPH,
                    order=order,
                    text=_text(item.text),
                    style=style,
                    is_bold=any(run.bold for run in item.runs),
                    # ponytail: inferred headings are all level 1. Font size
                    # does rank them, but Stage 2 gives a heading everything
                    # down to the next one of its level, so a ranked outline
                    # would hand a chapter its sub-sections' tables. Rank them
                    # once Stage 2 reads nested sections.
                    heading_level=1 if order in inferred else _heading_level(style),
                )
            )
        else:
            blocks.append(
                Table(
                    kind=BlockKind.TABLE,
                    order=order,
                    rows=[[_text(cell.text) for cell in row.cells] for row in item.rows],
                )
            )

    return Stage1Document(source=path.name, blocks=blocks)
