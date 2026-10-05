"""Agent 5: QA Adjudicator.

Agent 4 says what happened. This says what it means and who should act.

Why this is a separate step
---------------------------
A failing test is not a defect. On the first full S01 run, seven tests did not
pass and only two were the survey's fault; the rest were faults in the bot.
Had that run been reported as "seven defects", someone would have spent a day
investigating a survey that was behaving correctly.

Telling those apart needs the whole run rather than one test. Ten tests failing
on one question means something different from ten failing across ten
questions, and neither is visible from inside a single result.

What it decides
---------------
Every test that did not pass is given one cause and one owner:

    SURVEY_DEFECT       the built survey does not do what the questionnaire says
    SPECIFICATION_ERROR the questionnaire was read wrongly upstream
    NOT_BUILT_YET       the questionnaire asks for something not yet emitted
    UNSETTLED_QUESTION  the questionnaire never said, and a reading was assumed
    TEST_MODEL_GAP      the rule is enforced, but not in the way the test expects
    HARNESS_FAULT       the bot could not carry out the test

The last two matter. TEST_MODEL_GAP came out of a real run: a maximum length
enforced by preventing the respondent typing past it, rather than by rejecting
what they typed. The rule works. The test has no category for it. Reporting
that as a survey defect would be wrong, and hiding it would be worse, because
it is a real thing to fix in the test design.

The path view
-------------
A survey programmer does not think in test cases. They think in journeys: can a
respondent get through, and where do they come unstuck. So the report leads
with the paths, says which completed and which broke, and names the step it
broke at.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field, asdict
from datetime import datetime
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------- causes

SURVEY_DEFECT = "SURVEY_DEFECT"
SPECIFICATION_ERROR = "SPECIFICATION_ERROR"
NOT_BUILT_YET = "NOT_BUILT_YET"
UNSETTLED_QUESTION = "UNSETTLED_QUESTION"
TEST_MODEL_GAP = "TEST_MODEL_GAP"
HARNESS_FAULT = "HARNESS_FAULT"
COVERAGE_GAP = "COVERAGE_GAP"
NEEDS_SAMPLE = "NEEDS_SAMPLE"
UNDECIDED = "UNDECIDED"

OWNER = {
    SURVEY_DEFECT: "Survey Builder",
    SPECIFICATION_ERROR: "QRE Interpreter",
    NOT_BUILT_YET: "Survey Builder (already on the backlog)",
    UNSETTLED_QUESTION: "the client",
    TEST_MODEL_GAP: "Test Designer",
    HARNESS_FAULT: "Respondent Bot",
    COVERAGE_GAP: "Respondent Bot (not built yet)",
    NEEDS_SAMPLE: "nobody yet (needs a run with many respondents)",
    UNDECIDED: "needs a person to look",
}

MEANS = {
    SURVEY_DEFECT:
        "The built survey does not behave the way the questionnaire says it "
        "should. A respondent would see the wrong thing.",
    SPECIFICATION_ERROR:
        "The survey may be fine. The questionnaire was read wrongly further "
        "up the pipeline, so the test is asking for the wrong behaviour.",
    NOT_BUILT_YET:
        "The questionnaire asks for something the survey builder does not "
        "emit yet. Expected to fail until it does.",
    UNSETTLED_QUESTION:
        "The questionnaire never settled this, and a reading was assumed. "
        "Neither the survey nor the test is wrong until someone decides.",
    TEST_MODEL_GAP:
        "The rule is enforced, but by a different mechanism than the test "
        "looks for. Nothing is broken; the test needs to learn the other way.",
    HARNESS_FAULT:
        "The bot could not carry the test out, so nothing was proved either "
        "way. This says nothing about the survey.",
    COVERAGE_GAP:
        "The bot reached the right place but does not yet know how to observe "
        "this kind of claim, so nothing was proved either way. Nothing is "
        "known to be wrong.",
    NEEDS_SAMPLE:
        "This can only be checked across many respondents, such as the order "
        "options are shown in. One run through the survey cannot prove it, so "
        "it was never going to run as a single journey.",
    UNDECIDED:
        "No rule matched this one. A person should read it.",
}

# Ranked by whether a real respondent's data would be wrong.
SEVERITY = {
    SURVEY_DEFECT: 1,
    SPECIFICATION_ERROR: 2,
    NOT_BUILT_YET: 3,
    TEST_MODEL_GAP: 4,
    UNSETTLED_QUESTION: 5,
    HARNESS_FAULT: 6,
    COVERAGE_GAP: 7,
    NEEDS_SAMPLE: 8,
    UNDECIDED: 9,
}


@dataclass
class Judgement:
    test_case_id: str
    question: str
    path_id: str
    title: str
    outcome: str
    cause: str
    owner: str
    reason: str                 # why this cause and not another
    to_fix: str                 # what someone should actually do
    observed: str = ""
    order: int = 0


@dataclass
class PathVerdict:
    path_id: str
    name: str
    route_class: str
    tests: int = 0
    passed: int = 0
    failed: int = 0
    other: int = 0
    broke_at: str = ""          # the earliest question that failed
    first_failure: str = ""
    first_order: int = 9999
    real_broke_at: str = ""     # the earliest failure a respondent would feel
    real_first_failure: str = ""
    real_first_order: int = 9999
    causes: list = field(default_factory=list)

    #: Causes a respondent would actually notice. A journey is only broken if
    #: something on it would go wrong for a real person. A gap in our own test
    #: design failing on a journey says nothing about that journey, and
    #: calling it broken sends the reader to the wrong place, which is the
    #: first thing they read.
    _REAL = (SURVEY_DEFECT, SPECIFICATION_ERROR, NOT_BUILT_YET)

    @property
    def status(self) -> str:
        if any(c in self._REAL for c in self.causes):
            return "BROKEN"
        if self.failed or self.other:
            return "NOT PROVEN"
        return "WORKING"

    @property
    def status_note(self) -> str:
        if self.status == "WORKING":
            return "every test on this journey passed"
        if self.status == "BROKEN":
            return ("a respondent taking this journey would meet something "
                    "wrong")
        return ("nothing on this journey is known to be wrong, but not "
                "everything could be proved: "
                + ", ".join(self.causes))


# ---------------------------------------------------------------- deciding

_RULE_ID = re.compile(r"\bR\d+\b", re.I)


def _subjects(finding: dict) -> set[str]:
    """What a finding is about: questions, rules, quota names."""
    return {part.strip().split("/")[0]
            for part in (finding.get("subject") or "").split(",")
            if part.strip()}


def _explains(finding: dict, question: str, text: str) -> bool:
    """Whether one finding actually accounts for this test.

    A finding about a rule names the rule (R19), a finding about a question
    names the question, so each is matched on what it names. Matching a rule
    finding against a question never worked: "R20" is not "Q6", so the survey
    was blamed for something Agent 3 had already said was not built.

    A finding about a question must also be about what the test examines.
    Randomization not being built on Q1 explains a test about the order of
    Q1's options. It says nothing about whether Q1 refuses an exclusive
    option, and blaming it for that sends the reader to the wrong place.
    """
    kind = finding.get("kind")
    subjects = _subjects(finding)
    low = text.lower()
    if kind in ("SKIP_RULE_NOT_BUILT", "REJECT_RULE_NOT_BUILT"):
        return bool(subjects & {m.upper() for m in _RULE_ID.findall(text)})
    if kind == "RANDOMIZATION_NOT_BUILT":
        return question in subjects and ("order" in low or "random" in low)
    if kind == "QUOTAS_NOT_BUILT":
        return any(name.lower() in low for name in subjects)
    return question in subjects


def _decide(result: dict, index_row: dict, conformance: list[dict],
            provisional: list[str],
            built: dict | None = None) -> tuple[str, str, str]:
    """One test's cause, why, and what to do about it.

    The order of these rules is the argument. A failure the bot caused is not
    evidence about the survey, so that is asked first. A behaviour the builder
    has not written yet cannot be a defect in it, so that is asked next. Only
    when nothing else explains the failure is the survey blamed.
    """
    status = result.get("status")
    question = index_row.get("question", "")
    title = (result.get("title") or "").lower()
    seen = " ".join(str(c.get("actually", "")) for c in result.get("checks", []))

    # 1. Did the test actually run
    if status == "SKIPPED" and "several respondents" in str(
            result.get("blocked_reason") or ""):
        return (NEEDS_SAMPLE,
                "the claim is about what happens across many respondents, so "
                "it is not a single journey and the bot rightly did not run it",
                "no action on this test. It needs a run with a sample, which "
                "does not exist yet")

    if status in ("BLOCKED", "SKIPPED"):
        return (HARNESS_FAULT,
                "the journey never reached the point it was testing, so "
                "nothing about the survey was observed",
                "fix the bot, then run this test again. Until then it is "
                "neither a pass nor a failure")

    if status == "INCONCLUSIVE" and "cannot observe" in seen:
        return (COVERAGE_GAP,
                "the bot reached the right place but has no way yet to observe "
                "this kind of claim",
                "teach the bot to observe it, or record that this claim cannot "
                "be proved by running one respondent through the survey")

    if status == "INCONCLUSIVE":
        return (TEST_MODEL_GAP,
                "the bot reached the right place but could not tell whether "
                "the claim held",
                "give the test something observable to check, or record that "
                "this claim cannot be proved by running the survey")

    named = f"{result.get('title') or ''} {index_row.get('test_name') or ''}"

    # 2. Is the behaviour even built yet
    for f in conformance:
        if not _explains(f, question, named):
            continue
        if f.get("kind") in ("RANDOMIZATION_NOT_BUILT", "QUOTAS_NOT_BUILT",
                             "SKIP_RULE_NOT_BUILT", "REJECT_RULE_NOT_BUILT",
                             "GUARD_NOT_BUILT"):
            return (NOT_BUILT_YET,
                    f"before the run, Agent 3 already reported "
                    f"{f['kind']} for {f['subject']}",
                    "no action on this test. It will pass once the survey "
                    "builder emits that rule")

    # 3. Was the questionnaire read wrongly
    for f in conformance:
        if not _explains(f, question, named):
            continue
        if f.get("kind") in ("GUARD_CANNOT_BE_FALSIFIED",
                             "DISPOSITION_NOT_DISTINGUISHABLE"):
            return (SPECIFICATION_ERROR,
                    f"Agent 3 already reported {f['kind']} here: "
                    f"{f.get('detail','')[:110]}",
                    "correct the extraction upstream and regenerate. The "
                    "survey may well be fine")

    # 3b. An exclusive option enforced by LimeSurvey unticking the others
    attrs = (built or {}).get(question) or {}
    if ("exclusive_option_id" in title and attrs.get("exclude_all_others")
            and "accepted the answer and moved on" in seen):
        return (TEST_MODEL_GAP,
                "the survey was built with LimeSurvey's exclusive option "
                "(exclude_all_others), which unticks the other choices when "
                "the exclusive one is ticked. The respondent cannot submit "
                "both, so there is no error to show",
                "check by hand that ticking the exclusive option clears the "
                "others. If it does, the rule works and the test should check "
                "the end state, not an error message")

    # 4. A rule enforced a different way than the test expects
    if ("rejects an answer" in title or "refuses" in title) and \
            "accepted the answer and moved on" in seen and \
            "no error message appeared" in seen:
        return (TEST_MODEL_GAP,
                "the survey accepted the answer silently and showed no error, "
                "which is what happens when a limit is enforced by stopping "
                "the respondent entering it rather than by rejecting what "
                "they entered",
                "check how the rule is built. If the respondent cannot break "
                "it in the first place, the rule works and the test should "
                "say so differently")

    # 5. Does it rest on something nobody has ruled on
    if provisional:
        return (UNSETTLED_QUESTION,
                f"this test depends on an unconfirmed reading: "
                f"{', '.join(provisional)}",
                "ask the client which reading is intended. Until then this is "
                "a question, not a defect")

    # 6. Nothing else explains it
    if status == "FAILED":
        return (SURVEY_DEFECT,
                "the survey was reachable, the test ran, and the behaviour "
                "the questionnaire describes did not happen",
                f"check how {question or 'this question'} is built in "
                f"LimeSurvey against what the questionnaire says")

    return (UNDECIDED, "no rule matched", "read this one by hand")


# ---------------------------------------------------------------- the run

def adjudicate(directory: Path) -> dict:
    a4 = json.loads((directory / "agent4" / "agent4_results.json")
                    .read_text(encoding="utf-8"))
    index = {r["executable_id"]: r for r in json.loads(
        (directory / "agent3" / "agent3_test_case_index.json")
        .read_text(encoding="utf-8"))["tests"] if r.get("executable_id")}

    conf_file = directory / "agent3" / "agent3_conformance.json"
    conformance = []
    if conf_file.exists():
        conformance = (json.loads(conf_file.read_text(encoding="utf-8"))
                       .get("content", {}).get("findings", []))

    # What the survey was actually built with, read from the file Agent 3
    # derived from the .lss. Lets a rule check the mechanism, not guess at it.
    built: dict = {}
    snap_file = directory / "agent3" / "implementation_snapshot.json"
    if snap_file.exists():
        snap = json.loads(snap_file.read_text(encoding="utf-8"))
        built = {qid: (q.get("attributes") or {})
                 for qid, q in (snap.get("content") or snap)
                 .get("questions", {}).items()}

    paths = []
    paths_file = directory / "agent3" / "agent3_paths.json"
    if paths_file.exists():
        paths = (json.loads(paths_file.read_text(encoding="utf-8"))
                 .get("content", {}).get("paths", []))

    judgements: list[Judgement] = []
    verdicts: dict[str, PathVerdict] = {
        p["path_id"]: PathVerdict(p["path_id"], p.get("name", ""),
                                  p.get("route_class", ""))
        for p in paths}

    for r in a4.get("results", []):
        if not r.get("status"):
            # A row with no outcome is not a result. Counting it would inflate
            # the totals with something that never ran.
            continue
        row = index.get(r.get("test_id"), {})
        pid = row.get("path_id") or ""
        v = verdicts.get(pid)
        if v:
            v.tests += 1

        if r.get("status") == "PASSED":
            if v:
                v.passed += 1
            continue

        cause, reason, to_fix = _decide(
            r, row, conformance, r.get("provisional") or [], built)
        j = Judgement(
            test_case_id=row.get("test_case_id") or r.get("test_id", ""),
            question=row.get("question", ""),
            path_id=pid,
            title=r.get("title", ""),
            outcome=r.get("status", ""),
            cause=cause, owner=OWNER[cause], reason=reason, to_fix=to_fix,
            order=row.get("order", 0),
            observed="; ".join(str(c.get("actually", ""))
                               for c in r.get("checks", [])
                               if c.get("matched") is False)
                     or r.get("blocked_reason", ""))
        judgements.append(j)

        if v:
            if r.get("status") == "FAILED":
                v.failed += 1
            else:
                v.other += 1
            if cause not in v.causes:
                v.causes.append(cause)
            # The earliest question in questionnaire order is where the
            # journey first came unstuck, which is the one worth naming.
            if j.order and j.order < v.first_order:
                v.broke_at = j.question or "start"
                v.first_failure = j.title
                v.first_order = j.order
            if (cause in PathVerdict._REAL and j.order
                    and j.order < v.real_first_order):
                v.real_broke_at = j.question or "start"
                v.real_first_failure = j.title
                v.real_first_order = j.order

    # A broken journey should name where it really broke. Otherwise an
    # earlier test the bot merely could not observe gets the blame, and the
    # reader goes looking at a question that is fine.
    for v in verdicts.values():
        if v.status == "BROKEN" and v.real_first_order < 9999:
            v.broke_at = v.real_broke_at
            v.first_failure = v.real_first_failure
            v.first_order = v.real_first_order

    return {"judgements": judgements, "paths": verdicts, "a4": a4}


# ---------------------------------------------------------------- grouping

def group(judgements: list[Judgement]) -> list[dict]:
    """Collect failures that share a cause and a question.

    One mistake upstream can fail nine tests. Listing nine defects invites nine
    investigations of one problem, so they are reported once with the nine
    underneath.
    """
    buckets: dict[tuple, list[Judgement]] = defaultdict(list)
    for j in judgements:
        buckets[(j.cause, j.question)].append(j)

    out = []
    for (cause, question), group_js in buckets.items():
        out.append({
            "cause": cause,
            "owner": OWNER[cause],
            "question": question or "whole survey",
            "count": len(group_js),
            "means": MEANS[cause],
            "reason": group_js[0].reason,
            "to_fix": group_js[0].to_fix,
            "tests": [j.test_case_id or "(unnamed)" for j in group_js],
            "severity": SEVERITY[cause],
            "order": min((j.order for j in group_js), default=999),
        })
    out.sort(key=lambda g: (g["severity"], g["order"]))
    return out


# ---------------------------------------------------------------- the report

PATH_FILL = {"WORKING": "DAEDEA", "BROKEN": "F8E4E6",
             "NOT PROVEN": "FAEBDE"}
PATH_FONT = {"WORKING": "0F766E", "BROKEN": "A32C36",
             "NOT PROVEN": "9C4709"}
CAUSE_FILL = {
    SURVEY_DEFECT: "F8E4E6", SPECIFICATION_ERROR: "FAEBDE",
    NOT_BUILT_YET: "EFEFEF", UNSETTLED_QUESTION: "E4E1F7",
    TEST_MODEL_GAP: "E4E1F7", HARNESS_FAULT: "EFEFEF", UNDECIDED: "FAEBDE",
    COVERAGE_GAP: "EFEFEF", NEEDS_SAMPLE: "EFEFEF",
}
CAUSE_FONT = {
    SURVEY_DEFECT: "A32C36", SPECIFICATION_ERROR: "9C4709",
    NOT_BUILT_YET: "635E7E", UNSETTLED_QUESTION: "4338A8",
    TEST_MODEL_GAP: "4338A8", HARNESS_FAULT: "635E7E", UNDECIDED: "9C4709",
    COVERAGE_GAP: "635E7E", NEEDS_SAMPLE: "635E7E",
}
INK = "1B1832"


def write_report(path: Path, data: dict, survey: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="C7C2DC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    judgements = data["judgements"]
    verdicts = list(data["paths"].values())
    groups = group(judgements)
    counts = data["a4"].get("counts", {})

    wb = Workbook()

    # ---- 1. the journeys, because that is how a programmer thinks --------
    ws = wb.active
    ws.title = "Respondent journeys"
    ws["A1"] = f"Which journeys through {survey} work"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=INK)
    ws["A2"] = ("A journey is one complete route a respondent can take. Each "
                "one carries the tests that can only be checked along it, so "
                "a broken journey means a respondent taking that route meets "
                "something wrong.")
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="4A4566")
    ws.merge_cells("A2:H2")

    head = ["Path", "The journey", "Kind", "Status", "What the status means",
            "Tests on it", "Passed", "Came unstuck at",
            "What went wrong first"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(vertical="bottom", wrap_text=True)
    for col, w in zip("ABCDEFGHI", (8, 40, 19, 14, 44, 11, 9, 16, 44)):
        ws.column_dimensions[col].width = w

    for r, v in enumerate(sorted(verdicts, key=lambda x: x.path_id), start=5):
        values = [v.path_id, v.name, v.route_class, v.status, v.status_note,
                  v.tests, v.passed, v.broke_at or "-",
                  v.first_failure or "-"]
        for i, val in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = box
            c.alignment = wrap
        ws.cell(row=r, column=4).font = Font(
            name="Calibri", size=10, bold=True, color=PATH_FONT[v.status])
        ws.cell(row=r, column=4).fill = PatternFill(
            "solid", fgColor=PATH_FILL[v.status])
    ws.freeze_panes = "A5"

    note = ws.cell(row=len(verdicts) + 6, column=1, value=(
        "WORKING means every test on this journey passed.\n"
        "BROKEN means a respondent taking it would meet something wrong: a "
        "defect in the survey, a misreading of the questionnaire, or a rule "
        "not built yet.\n"
        "NOT PROVEN means nothing is known to be wrong, but something could "
        "not be checked. A gap in our own test design failing on a journey "
        "says nothing about that journey, so it is not called broken."))
    note.font = Font(name="Calibri", size=10, italic=True, color="4A4566")
    note.alignment = wrap
    ws.merge_cells(start_row=len(verdicts) + 6, start_column=1,
                   end_row=len(verdicts) + 6, end_column=9)
    ws.row_dimensions[len(verdicts) + 6].height = 44

    # ---- 2. what to fix, grouped -----------------------------------------
    ws = wb.create_sheet("What to fix")
    ws["A1"] = "What needs correcting, and by whom"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=INK)
    ws["A2"] = ("Failures that share a cause and a question are reported once. "
                "One mistake upstream can fail nine tests, and nine entries "
                "would invite nine investigations of one problem. Ordered so "
                "that anything a real respondent would notice comes first.")
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="4A4566")
    ws.merge_cells("A2:G2")

    head = ["Cause", "Question", "Tests", "Who should act",
            "What this means", "Why it was judged this way", "What to do"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(vertical="bottom", wrap_text=True)
    for col, w in zip("ABCDEFG", (22, 11, 26, 24, 46, 50, 50)):
        ws.column_dimensions[col].width = w

    for r, g in enumerate(groups, start=5):
        values = [g["cause"], g["question"], ", ".join(g["tests"]),
                  g["owner"], g["means"], g["reason"], g["to_fix"]]
        for i, val in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = box
            c.alignment = wrap
        ws.cell(row=r, column=1).font = Font(
            name="Calibri", size=10, bold=True, color=CAUSE_FONT[g["cause"]])
        ws.cell(row=r, column=1).fill = PatternFill(
            "solid", fgColor=CAUSE_FILL[g["cause"]])
    ws.freeze_panes = "A5"

    # ---- 3. every judgement, for anyone who wants the detail -------------
    ws = wb.create_sheet("Every test judged")
    head = ["Order", "Test Case ID", "Question", "Path", "What was tested",
            "Outcome", "Cause", "What the bot saw", "What to do"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(vertical="bottom", wrap_text=True)
    for col, w in zip("ABCDEFGHI", (7, 13, 10, 7, 44, 13, 22, 48, 46)):
        ws.column_dimensions[col].width = w

    for r, j in enumerate(sorted(judgements, key=lambda x: x.order), start=2):
        values = [j.order, j.test_case_id, j.question, j.path_id, j.title,
                  j.outcome, j.cause, j.observed, j.to_fix]
        for i, val in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name="Calibri", size=10)
            c.border = box
            c.alignment = wrap
        ws.cell(row=r, column=7).font = Font(
            name="Calibri", size=10, bold=True, color=CAUSE_FONT[j.cause])
        ws.cell(row=r, column=7).fill = PatternFill(
            "solid", fgColor=CAUSE_FILL[j.cause])
    ws.freeze_panes = "A2"
    if judgements:
        ws.auto_filter.ref = f"A1:I{len(judgements) + 1}"

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Agent 5: adjudicate a run.")
    ap.add_argument("survey_dir")
    args = ap.parse_args()

    directory = Path(args.survey_dir)
    for needed in ("agent4/agent4_results.json",
                   "agent3/agent3_test_case_index.json"):
        if not (directory / needed).exists():
            print(f"missing {needed}. Run Agent 3 and then Agent 4 first.")
            return 1

    data = adjudicate(directory)
    groups = group(data["judgements"])
    verdicts = list(data["paths"].values())

    print(f"\n  Agent 5: {directory.name}")
    print("  " + "=" * 62)
    working = sum(1 for v in verdicts if v.status == "WORKING")
    print(f"  journeys      {working} of {len(verdicts)} working")
    passed = data["a4"].get("counts", {}).get("PASSED", 0)
    total = sum(data["a4"].get("counts", {}).values())
    print(f"  tests         {passed} of {total} passed")
    print()
    for v in sorted(verdicts, key=lambda x: x.path_id):
        mark = {"WORKING": "ok  ", "BROKEN": "FAIL", "NOT PROVEN": "?   "}
        print(f"    {mark[v.status]}  {v.path_id}  {v.name[:46]:<48}"
              + (f"broke at {v.broke_at}" if v.broke_at else ""))
    print()
    if groups:
        print("  WHAT TO FIX")
        for g in groups:
            print(f"    {g['cause']:<20} {g['question']:<6} "
                  f"{g['count']} test(s)   -> {g['owner']}")
            print(f"        {g['to_fix'][:92]}")
    else:
        print("  Nothing to fix. Every test passed.")

    dest = directory / "agent5"
    dest.mkdir(exist_ok=True)
    book = dest / "agent5_qc_report.xlsx"
    try:
        write_report(book, data, directory.name)
    except PermissionError:
        book = dest / f"agent5_qc_report_{datetime.now():%H%M%S}.xlsx"
        write_report(book, data, directory.name)
    (dest / "agent5_findings.json").write_text(json.dumps({
        "survey": directory.name,
        "run_at": datetime.now().isoformat(timespec="seconds"),
        # status and status_note are properties, which asdict leaves out.
        # The dashboard reads them, so they are written explicitly.
        "journeys": [{**asdict(v), "status": v.status,
                      "status_note": v.status_note} for v in verdicts],
        "groups": groups,
        "judgements": [asdict(j) for j in data["judgements"]],
    }, indent=2, default=str), encoding="utf-8")
    print(f"\n  wrote {book}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
