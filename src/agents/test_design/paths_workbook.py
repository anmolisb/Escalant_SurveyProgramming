"""The reviewable workbook, in the v5 template shape.

Three tabs:

  Survey Paths   The exclusive, branch-exhaustive set of respondent journeys.
                 A testing team starts here, because a journey is the unit a
                 person can picture and sign off. Everything after hangs off it.
  Test Cases     One row per test, in the v5 column set, tagged with the path
                 it runs on and the question it belongs to, ordered the way the
                 questionnaire is ordered.
  Not Enumerated What was deliberately left out of the path set, how large it
                 was, and what risk remains.

Two additions to the v5 column set, both requested:

  Question       Which questionnaire item the test hangs off, so a reader can
                 see at a glance what a test is about.
  Order          A running number in questionnaire order, S1 first, so the
                 sheet reads the same way the Word file does.
"""

from __future__ import annotations

from pathlib import Path as FsPath

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .a_implementation import ImplementationSnapshot
from .models import COVERED, CoverageTarget, VerifiedScenario
from .path_enumerator import Path, Branch, host_path
from .spec import CanonicalSpec

F = "Calibri"
HEAD_FILL = PatternFill("solid", fgColor="1B1832")
SUB_FILL = PatternFill("solid", fgColor="E9E7F1")
HEAD_F = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE_F = Font(name=F, size=14, bold=True, color="1B1832")
LEAD_F = Font(name=F, size=10, italic=True, color="4A4566")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
MONO = Font(name="Consolas", size=9)
OK = Font(name=F, size=10, bold=True, color="0F766E")
WARN = Font(name=F, size=10, bold=True, color="9C4709")
THIN = Side(style="thin", color="C7C2DC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")


# --------------------------------------------------------------------------
# Vocabulary
# --------------------------------------------------------------------------

# The v5 test classes.
CLS_ROUTE = "ROUTE"
CLS_SEGMENT = "SEGMENT"
CLS_FOCUSED = "FOCUSED"
CLS_STATEFUL = "STATEFUL"
CLS_PROBE = "PROBE"
CLS_STATIC = "STATIC"

# The v5 test types, keyed by our internal dimension.
PRIMARY_TYPE = {
    "D1": "1. Flow / Routing",
    "D2": "6. Termination / Disposition",
    "D3": "2. Input Validation",
    "D4": "2. Input Validation",
    "D5": "3. Dependency / Piping",
    "D6": "3. Dependency / Piping",
    "D7": "4. Randomization",
    "D8": "5. Quota / Stateful",
    "D9": "9. Rule Interaction",
}

SECONDARY_TYPE = {
    ("D1", "advances"): "-",
    ("D1", "shown"): "-",
    ("D1", "hidden"): "10. Negative / Robustness",
    ("D1", "skip_fired"): "-",
    ("D1", "skip_not_fired"): "10. Negative / Robustness",
    ("D2", "reachable"): "1. Flow / Routing",
    ("D3", "satisfied"): "-",
    ("D3", "violated"): "10. Negative / Robustness",
    ("D3", "boundary_max_accepted"): "10. Negative / Robustness",
    ("D3", "special_characters_accepted"): "10. Negative / Robustness",
    ("D4", "whitespace_rejected"): "10. Negative / Robustness",
    ("D4", "whitespace_accepted"): "-",
    ("D4", "enforced"): "10. Negative / Robustness",
    ("D4", "not_enforced"): "-",
    ("D5", "restricted"): "1. Flow / Routing",
    ("D6", "rendered"): "-",
    ("D7", "completeness"): "8. Implementation Conformance",
    ("D7", "order_varies"): "8. Implementation Conformance",
    ("D7", "anchors_held"): "8. Implementation Conformance",
    ("D8", "available"): "6. Termination / Disposition",
    ("D8", "full"): "6. Termination / Disposition",
    ("D8", "over_target_admits"): "5. Quota / Stateful",
    ("D8", "not_counted_by_other_cell"): "5. Quota / Stateful",
    ("D9", "combined"): "1. Flow / Routing",
}

PRIORITY = {
    "D1": "High", "D2": "Critical", "D3": "High", "D4": "Medium",
    "D5": "High", "D6": "Medium", "D7": "Medium", "D8": "Critical",
    "D9": "Medium",
}

WHY_EXISTS = {
    "D1": "Routing is where survey builds most often go wrong. A question that "
          "appears for the wrong people, or does not appear for the right ones, "
          "corrupts the data without anyone noticing.",
    "D2": "Screening someone out incorrectly loses a valid respondent and the "
          "money spent recruiting them. Failing to screen someone out puts the "
          "wrong person in the sample.",
    "D3": "A rule that is written down but not enforced lets unusable answers "
          "into the data. One that is over-enforced blocks valid respondents.",
    "D4": "A compulsory question that lets people through blank produces "
          "missing data that cannot be recovered afterwards.",
    "D5": "A question that quietly shows the wrong options asks the respondent "
          "about something they never selected.",
    "D6": "A question that quotes the wrong earlier answer reads as a different "
          "question entirely to the respondent.",
    "D7": "Without a genuine shuffle the option order biases the answers, and "
          "without the recorded order the answers cannot be corrected later.",
    "D8": "Quotas decide who gets turned away. Getting this wrong unbalances "
          "the whole sample.",
    "D9": "Individually correct rules can combine into behaviour neither one "
          "describes on its own.",
}

TEST_DATA_CLASS = {
    ("D3", "satisfied"): "VALID_BOUNDARY",
    ("D3", "violated"): "INVALID_BOUNDARY",
    ("D3", "boundary_max_accepted"): "VALID_BOUNDARY (exactly at the limit)",
    ("D3", "special_characters_accepted"): "VALID_SPECIAL_CHARACTERS",
    ("D4", "whitespace_rejected"): "WHITESPACE_ONLY",
    ("D4", "whitespace_accepted"): "WHITESPACE_ONLY",
    ("D4", "enforced"): "EMPTY",
    ("D4", "not_enforced"): "EMPTY",
    ("D1", "hidden"): "VALID_TYPICAL (the branch that hides the question)",
    ("D8", "full"): "VALID_TYPICAL, with the cell pre-filled",
}

SUB_ORDER = {
    ("D1", "advances"): 0, ("D1", "shown"): 1, ("D1", "hidden"): 2,
    ("D1", "skip_fired"): 3, ("D1", "skip_not_fired"): 4,
    ("D3", "satisfied"): 5, ("D3", "violated"): 6,
    ("D3", "boundary_max_accepted"): 6, ("D3", "special_characters_accepted"): 6,
    ("D4", "whitespace_rejected"): 7, ("D4", "whitespace_accepted"): 7,
    ("D4", "enforced"): 7, ("D4", "not_enforced"): 7,
    ("D5", "restricted"): 8, ("D6", "rendered"): 9,
    ("D7", "completeness"): 10, ("D7", "order_varies"): 11,
    ("D7", "anchors_held"): 12,
    ("D8", "available"): 13, ("D8", "full"): 14,
    ("D8", "over_target_admits"): 14, ("D8", "not_counted_by_other_cell"): 15,
    ("D9", "combined"): 15, ("D2", "reachable"): 16,
}

COLUMNS = [
    ("Order", 7), ("Question", 11), ("Test Case ID", 14), ("Test Class", 11),
    ("Primary Test Type", 24), ("Secondary Test Types", 24), ("Path ID", 9),
    ("Test Name", 46), ("Test Objective", 54), ("Why This Test Exists", 50),
    ("Priority", 9), ("Preconditions", 34), ("Isolation Mode", 17),
    ("Respondent Inputs", 44), ("Test Data Class", 20),
    ("Why Inputs Were Chosen", 46), ("Actions", 46),
    ("Expected Question Sequence", 40), ("Expected Questions Presented", 26),
    ("Expected Questions Skipped", 24), ("Expected Questions Hidden / Not Presented", 26),
    ("Expected Navigation / Next Question", 34), ("Expected Validation", 40),
    ("Expected Disposition", 20), ("Expected Evidence", 40),
    ("QRE Traceability", 46), ("Evidence Provenance", 34),
    ("Status", 11), ("Runnable", 10),
    ("What This Test Does NOT Prove", 46), ("Notes", 30),
]


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def _sheet(wb, first, title, lead, headers, widths):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws["A1"] = title
    ws["A1"].font = TITLE_F
    ws["A2"] = lead
    ws["A2"].font = LEAD_F
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=max(len(headers), 2))
    ws.row_dimensions[2].height = 30
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_F, BOX
        c.alignment = Alignment(vertical="bottom", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}4"
    ws._next = 5
    return ws


def _row(ws, values, fonts=None):
    r = ws._next
    ws._next += 1
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i,
                    value=(v if isinstance(v, (int, float, str)) or v is None
                           else str(v)))
        c.font = (fonts or {}).get(i, BODY)
        c.border = BOX
        c.alignment = WRAP if i > 2 else TOP
    return r


def _labels(spec: CanonicalSpec, qid: str, value) -> str:
    q = spec.question(qid)
    if value is None or value == "" or value == [] or value == {}:
        return "left blank"
    if isinstance(value, dict):
        parts = []
        for k, v in list(value.items())[:4]:
            ko = q.option_by_id(str(k)) if q else None
            vo = q.option_by_id(str(v)) if q else None
            parts.append(f"{(ko.label if ko else k)}={(vo.label if vo else v)}")
        return "; ".join(parts) + ("; ..." if len(value) > 4 else "")
    if isinstance(value, (list, tuple)):
        out = []
        for v in value:
            o = q.option_by_id(str(v)) if q else None
            out.append(o.label if o else str(v))
        return " + ".join(out)
    o = q.option_by_id(str(value)) if q else None
    return o.label if o else str(value)


CHECK_WORDS = {
    "field_present": "{q} is on screen",
    "field_absent": "{q} is NOT on screen",
    "question_visible": "{q} is on screen",
    "question_absent": "{q} is NOT on screen",
    "respondent_continues": "the respondent carries on rather than being "
                            "screened out or finished",
    "not_on_end_page": "the respondent carries on rather than being screened "
                       "out or finished",
    "next_question_is": "the next question shown is {v}",
    "next_field_present": "the next question shown is {v}",
    "page_advances": "the survey ACCEPTS the answer and moves on",
    "answer_accepted": "the survey ACCEPTS the answer and moves on",
    "page_does_not_advance": "the survey REFUSES to move on; the same page "
                             "reappears",
    "answer_rejected": "the survey REFUSES to move on; the same page reappears",
    "error_shown_on_question": "an error message is shown against {q}",
    "survey_completed": "the respondent reaches the normal thank-you page",
    "group_suppressed": "no main-section question appears at all",
    "end_page_message": "the final page shows: {v}",
    "end_page_message_non_discriminating":
        "the final page shows: {v}   [other screen-outs show identical wording, "
        "so this proves a screen-out happened, not which one]",
    "not_sent_to_quota_full": "the respondent is allowed to continue",
    "quota_counter_over_target": "the counter for {v} records this respondent, "
                                 "putting the cell over its target",
    "quota_counter_unchanged": "the counter for {v} is exactly what it was "
                               "before this respondent",
    "question_text_contains": "{q}'s wording includes: {v}",
    "options_rendered_equals": "{q} offers exactly: {v}",
    "rendered_option_count": "{q} shows exactly {v} options, none missing or "
                             "repeated",
    "rendered_option_labels": "the options shown at {q} are exactly: {v}",
    "rendered_order_varies": "the option order at {q} is NOT identical on every "
                             "run",
    "anchored_options_hold_position": "these options hold their position: {v}",
}

NEG = {"field_present": "field_absent", "question_visible": "question_absent"}


def _checks(assertions: list[dict]) -> str:
    out = []
    for a in assertions:
        kind = str(a.get("kind", ""))
        exp = a.get("expected")
        if exp is False and kind in NEG:
            kind = NEG[kind]
        tpl = CHECK_WORDS.get(kind, kind.replace("_", " "))
        tgt = str(a.get("target") or a.get("canonical") or "")
        val = "" if exp in (True, False, None) else (
            ", ".join(str(x) for x in exp) if isinstance(exp, list) else str(exp))
        out.append(tpl.replace("{q}", tgt).replace("{v}", val))
    return "\n".join(f"{i}. {x}" for i, x in enumerate(out, 1))


def _rule_words(q) -> str:
    v = q.validation.raw
    bits = []
    if v.get("min_selections") is not None:
        n = v["min_selections"]
        bits.append(f"at least {n} option{'s' if n != 1 else ''} must be selected")
    if v.get("min_length") is not None:
        bits.append(f"at least {v['min_length']} characters")
    if v.get("max_length") is not None:
        bits.append(f"at most {v['max_length']} characters")
    if v.get("sum_to") is not None:
        bits.append(f"the numbers must total exactly {v['sum_to']}")
    if v.get("require_each_row"):
        bits.append("every row must be answered")
    lbl = v.get("exclusive_option_label") or v.get("exclusive_option_id")
    if lbl:
        bits.append(f"\u201c{lbl}\u201d cannot be chosen alongside anything else")
    return "; ".join(bits) or ", ".join(q.validation.constraints)


# --------------------------------------------------------------------------
# Per-test narrative
# --------------------------------------------------------------------------

def _mechanism(spec: CanonicalSpec, t: CoverageTarget) -> str | None:
    """Which specific mechanism a test is about, where several share an ending."""
    if t.dimension == "D8":
        return t.subject.split(":")[0]
    if t.dimension == "D2":
        ending, _, named = t.subject.partition("<-")
        if named:
            return named
        quota = next((q for q in spec.quotas if q.on_full == ending), None)
        return quota.id if quota else None
    return None


def _anchor_question(spec: CanonicalSpec, t: CoverageTarget) -> str | None:
    d, subj = t.dimension, t.subject
    if d == "D1" and t.polarity in ("advances", "shown", "hidden"):
        return subj
    if d == "D1":
        rule = spec.rule(subj.split(":")[0])
        return rule.evaluation_point if rule else None
    if d == "D2":
        ending, _, named = subj.partition("<-")
        rule = spec.rule(named) if named else None
        if rule and rule.evaluation_point:
            return rule.evaluation_point
        for q in spec.quotas:
            if q.on_full == ending:
                return q.variable_question_id
        return None
    if d == "D3":
        if spec.question(subj):
            return subj
        rule = spec.rule(subj)
        return rule.evaluation_point if rule else None
    if d == "D4":
        return subj
    if d in ("D5", "D6"):
        return subj.split("->")[-1]
    if d == "D7":
        return subj
    if d == "D8":
        quota = next((x for x in spec.quotas if x.id == subj.split(":")[0]), None)
        return quota.variable_question_id if quota else None
    if d == "D9":
        return subj.split("+")[-1]
    return None


def _test_class(t: CoverageTarget) -> str:
    if t.dimension == "D2":
        return CLS_ROUTE
    if t.dimension == "D8" and t.polarity == "full":
        return CLS_STATEFUL
    if t.dimension == "D7" and t.polarity == "order_varies":
        return CLS_STATEFUL
    return CLS_FOCUSED


def _name(spec: CanonicalSpec, t: CoverageTarget) -> str:
    d, s, pol = t.dimension, t.subject, t.polarity
    if d == "D1" and pol == "advances":
        return f"{s} accepts a normal answer and the survey moves on"
    if d == "D1" and pol == "shown":
        return f"{s} is shown when its condition is met"
    if d == "D1" and pol == "hidden":
        return f"{s} is not shown when its condition is not met"
    if d == "D1" and pol == "skip_fired":
        return f"{s.split(':')[0]} jumps forward when its condition is met"
    if d == "D1":
        return f"{s.split(':')[0]} does not jump when its condition is not met"
    if d == "D2":
        ending, _, named = s.partition("<-")
        return f"End-to-end journey: {ending} by way of {named or 'routing'}"
    if d == "D3" and pol == "boundary_max_accepted":
        return f"{s} accepts an answer sitting exactly on its stated maximum"
    if d == "D3" and pol == "special_characters_accepted":
        return f"{s} accepts punctuation and quotation marks"
    if d == "D4" and pol == "whitespace_rejected":
        return f"{s} refuses an answer of spaces alone"
    if d == "D4" and pol == "whitespace_accepted":
        return f"{s} accepts spaces alone, being optional"
    if d == "D3":
        q = spec.question(s)
        if q is not None:
            return (f"{s} {'accepts a valid answer' if pol == 'satisfied' else 'rejects an invalid answer'}")
        return f"Reject rule {s} {'does not fire' if pol == 'satisfied' else 'fires'}"
    if d == "D4":
        return (f"{s} blocks a blank answer" if pol == "enforced"
                else f"{s} allows a blank answer, being optional")
    if d == "D5":
        a, b = s.split("->")
        return f"{b} offers only the options chosen at {a}"
    if d == "D6":
        a, b = s.split("->")
        return f"{b} quotes the answer given at {a}"
    if d == "D7":
        return {"completeness": f"{s} shuffles without losing an option",
                "order_varies": f"{s} genuinely varies its option order",
                "anchors_held": f"{s} keeps its pinned options in place"}[pol]
    if d == "D8":
        quota, cell = s.split(":")
        return {
            "available": f"{quota} admits a respondent while {cell} is open",
            "full": f"{quota} turns a respondent away once {cell} is full",
            "over_target_admits": f"{quota} still admits a respondent once "
                                  f"{cell} is over target, being a soft quota",
            "not_counted_by_other_cell": f"{quota} does not increment {cell} for "
                                         f"a respondent in a different cell",
        }.get(pol, f"{quota} {pol} at {cell}")
    if d == "D9":
        a, b = s.split("+")
        return f"{b} is shown and reflects its dependency on {a} at the same time"
    return t.claim


def _objective(spec: CanonicalSpec, t: CoverageTarget) -> str:
    d, s, pol = t.dimension, t.subject, t.polarity
    if d == "D1" and pol == "advances":
        return (f"Prove that answering {s} normally moves the respondent on to "
                f"whatever the questionnaire says comes next, rather than "
                f"stopping or jumping elsewhere.")
    if d == "D1" and pol in ("shown", "hidden"):
        q = spec.question(s)
        cond = q.guard.render() if (q and q.guard) else ""
        verb = "IS shown" if pol == "shown" else "is NOT shown"
        return (f"Prove {s} {verb} when its display rule is "
                f"{'satisfied' if pol == 'shown' else 'not satisfied'}. "
                f"The rule is: {cond}")
    if d == "D1":
        rid = s.split(":")[0]
        rule = spec.rule(rid)
        cond = rule.when.render() if (rule and rule.when) else ""
        return (f"Prove {rid} {'jumps to ' + str(rule.destination_id) if pol == 'skip_fired' else 'does not jump and the respondent continues in sequence'}. "
                f"The rule is: {cond}")
    if d == "D2":
        ending, _, named = s.partition("<-")
        rule = spec.rule(named) if named else None
        cond = rule.when.render() if (rule and rule.when) else ""
        where = rule.evaluation_point if rule else "?"
        disp = spec.disposition(ending)
        msg = (disp.message or "")[:80] if disp else ""
        return (f"Prove that a respondent answering as listed is shown exactly "
                f"this sequence of questions, in this order, and ends at "
                f"{ending} by way of {named or 'the routing'}, checked at "
                f"{where}. Condition: {cond}. They should then see: \u201c{msg}\u201d")
    if d == "D3" and pol == "boundary_max_accepted":
        q = spec.question(s)
        hi = (q.validation.get("max_length") if q else None) or \
             (q.validation.get("max_selections") if q else None)
        unit = "characters" if (q and q.validation.get("max_length")) else "selections"
        return (f"Prove {s} accepts an answer of exactly {hi} {unit}. A maximum "
                f"is a limit, so the limit itself is inside the rule. One more "
                f"than this is already tested as a rejection; an off-by-one in "
                f"the build sits precisely between the two.")

    if d == "D3" and pol == "special_characters_accepted":
        return (f"Prove {s} accepts an answer containing apostrophes, quotation "
                f"marks and angle brackets. Its rule constrains length and says "
                f"nothing about content, so this must be accepted. Quotes and "
                f"brackets are where survey tools most often break, because "
                f"nobody wrote the rule down and so nobody tested it.")

    if d == "D4" and pol == "whitespace_rejected":
        return (f"Prove {s} refuses an answer of spaces alone. The "
                f"questionnaire marks it compulsory, and spaces are not an "
                f"answer. This is the commoner real-world failure than a blank, "
                f"because a respondent pressing the space bar looks like a "
                f"respondent who answered.")

    if d == "D4" and pol == "whitespace_accepted":
        return (f"Prove {s} accepts an answer of spaces alone. The "
                f"questionnaire marks it optional, so nothing about the "
                f"response may block progress.")

    if d == "D3":
        q = spec.question(s)
        if q is not None:
            return (f"Prove {s} {'accepts' if pol == 'satisfied' else 'refuses'} "
                    f"an answer that {'obeys' if pol == 'satisfied' else 'breaks'} "
                    f"its stated rule. The rule is: {_rule_words(q)}")
        rule = spec.rule(s)
        return (f"Prove reject rule {s} "
                f"{'does not fire when its condition is not met' if pol == 'satisfied' else 'fires when its condition is met'}. "
                f"The rule is: {rule.when.render() if (rule and rule.when) else ''}")
    if d == "D4":
        q = spec.question(s)
        if pol == "enforced":
            return (f"Prove {s} refuses to let the respondent past when it is "
                    f"left blank. The questionnaire marks it compulsory, and it "
                    f"is a {q.kind if q else '?'} question.")
        return (f"Prove {s} lets the respondent through with no answer, because "
                f"the questionnaire marks it optional.")
    if d in ("D5", "D6"):
        a, b = s.split("->")
        what = "option list" if d == "D5" else "wording"
        return (f"Prove {b}'s {what} is built from the answer given at {a}, so "
                f"the respondent is only asked about what they actually chose.")
    if d == "D7":
        q = spec.question(s)
        n = len(q.options) if q else 0
        return {"completeness": f"Prove {s} still presents all {n} of its options "
                                f"when shuffled, with none dropped and none repeated.",
                "order_varies": f"Prove {s} is genuinely shuffled: the option order "
                                f"must differ between respondents. If it never "
                                f"changes, shuffling was not switched on.",
                "anchors_held": f"Prove {s} keeps its pinned options in the same "
                                f"position while the rest move."}[pol]
    if d == "D8":
        quota, cell = s.split(":")
        q = next((x for x in spec.quotas if x.id == quota), None)
        c = next((x for x in (q.cells if q else []) if x.option_id == cell), None)
        label = c.option_label if c else cell
        target = c.target_count if c else "?"
        if pol == "available":
            return (f"Prove {quota} lets a {label!r} respondent through while "
                    f"that group is below its target.")
        if pol == "over_target_admits":
            return (f"Prove {quota} is enforced softly: once {label!r} reaches "
                    f"its target of {target}, the next respondent must still be "
                    f"let through and the overflow recorded. A soft quota "
                    f"reports an imbalance, it does not turn people away.")
        if pol == "not_counted_by_other_cell":
            others = [x.option_label for x in (q.cells if q else [])
                      if x.option_id != cell]
            return (f"Prove a respondent in a different cell does not count "
                    f"against {label!r}. Send one who answers "
                    f"{others[0] if others else 'another option'!r}, then "
                    f"confirm the {label!r} count is unchanged. Getting this "
                    f"wrong silently mis-fills every cell in the quota.")
        return (f"Prove {quota} turns a {label!r} respondent away once that "
                f"group reaches its target of {target}.")
    if d == "D9":
        a, b = s.split("+")
        return (f"Prove {b} is shown by its display rule and simultaneously "
                f"reflects its dependency on {a}. Each mechanism can be correct "
                f"alone and still disagree together.")
    return t.claim


def _not_proven(spec: CanonicalSpec, t: CoverageTarget) -> str:
    d, pol = t.dimension, t.polarity
    common = ("It does not prove anything about questions after the one under "
              "test, because the journey stops there.")
    if d == "D1" and pol == "advances":
        return ("It does not prove the answer was validated, only that the "
                "survey accepted it and moved to the right place. " + common)
    if d == "D1" and pol == "hidden":
        return ("It does not prove why the question was absent beyond the "
                "display rule; the check confirms the respondent was still "
                "travelling and had not been screened out. " + common)
    if d == "D2":
        return ("It does not prove any validation rule: every answer on this "
                "path is deliberately valid. It does not prove randomization or "
                "quota behaviour.")
    if d == "D3":
        return ("It proves one direction of one rule at one question. The "
                "opposite direction is a separate test. " + common)
    if d == "D4":
        return ("It proves only the compulsory setting on this question. "
                "Whether the flag is set correctly on other questions is "
                "checked separately, question by question.")
    if d == "D7" and pol == "completeness":
        return ("It does not prove the order varies, only that nothing was lost "
                "or repeated. Order is a separate test needing several runs.")
    if d == "D7" and pol == "order_varies":
        return ("It does not prove any particular order is correct, only that "
                "the order is not fixed.")
    if d == "D8":
        return ("It does not prove the counting is accurate, only that the "
                "boundary behaves as stated at the moment tested.")
    return common


def _provenance(spec: CanonicalSpec, t: CoverageTarget) -> str:
    q = spec.question(t.subject)
    origin = None
    if q is not None and q.guard is not None:
        origin = q.guard.origin
    rule = spec.rule(t.subject.split("<-")[-1]) if "<-" in t.subject else spec.rule(t.subject)
    if rule is not None and rule.when is not None:
        origin = rule.when.origin
    if origin == "inferred":
        return ("Model-inferred. The questionnaire wrote this as prose rather "
                "than in a table, so Agent 1 read it with a model and raised it "
                "for sign-off. Confirm before relying on this test.")
    if origin == "derived":
        return "Derived by rule from the questionnaire's own tables."
    return "QRE explicit - stated directly in the questionnaire."


def _traceability(spec: CanonicalSpec, t: CoverageTarget) -> str:
    bits = []
    q = spec.question(_anchor_question(spec, t) or "")
    if q is not None:
        ref = (q.source_reference or {}).get("text")
        if ref:
            bits.append(str(ref))
    rules = [r for r in t.traces_to if spec.rule(r)]
    if rules:
        bits.append("Rules exercised: " + ", ".join(sorted(set(rules))))
    return "\n".join(bits) or "-"


# --------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------

def build(out_path: FsPath, spec: CanonicalSpec, targets: list[CoverageTarget],
          scenarios: list[VerifiedScenario], logical_tests: list[dict],
          exec_tests: list[dict], snap: ImplementationSnapshot | None,
          summary: dict, paths: list[Path], branches: list[Branch],
          path_report: dict, conformance: dict | None = None) -> FsPath:

    sc_by = {s.target_id: s for s in scenarios}
    log_by = {t["target_id"]: t for t in logical_tests}
    exe_by = {t["target_id"]: t for t in exec_tests}
    seq_of = {q.id: q.seq for q in spec.questions}

    # Work out the ordering and the display ids ONCE, before either tab is
    # written. The paths tab needs to name the tests it hosts, and those names
    # have to be the ids a reader will actually find in the test tab.
    ordered = []
    for t in targets:
        lt = log_by.get(t.target_id)
        sc = sc_by.get(t.target_id)
        if lt is None or sc is None or sc.status != COVERED:
            continue
        aq = _anchor_question(spec, t)
        ordered.append((seq_of.get(aq or "", 10 ** 6),
                        SUB_ORDER.get((t.dimension, t.polarity), 99),
                        t.subject, t, lt, sc, aq))
    ordered.sort(key=lambda r: (r[0], r[1], r[2]))

    counters = {CLS_ROUTE: 0, CLS_FOCUSED: 0, CLS_STATEFUL: 0,
                CLS_SEGMENT: 0, CLS_PROBE: 0, CLS_STATIC: 0}
    prefix = {CLS_ROUTE: "RT", CLS_FOCUSED: "FT", CLS_STATEFUL: "ST",
              CLS_SEGMENT: "SG", CLS_PROBE: "PR", CLS_STATIC: "SC"}
    display_id: dict[str, str] = {}
    for _, _, _, t, _, _, _ in ordered:
        cls = _test_class(t)
        counters[cls] += 1
        display_id[t.target_id] = f"TC-{prefix[cls]}{counters[cls]:02d}"

    wb = Workbook()

    # =================== TAB 1: SURVEY PATHS ===============================
    ws = _sheet(
        wb, True, "Survey Paths",
        "The exclusive, branch-exhaustive set of respondent journeys through "
        "this survey. Exclusive means no two paths walk the same question "
        "sequence. Branch-exhaustive means every point where the survey can go "
        "more than one way is taken in both directions by at least one path. "
        "Start here: a journey is the unit a testing team can picture, run and "
        "sign off, and every test in the next tab is hosted on one of these.",
        ["Path ID", "Path Name", "Route Class", "Preconditions",
         "Branch Decisions", "Branch Values", "Expected Question Sequence",
         "Questions Presented", "Expected Skips", "Expected Hidden / Not Presented",
         "Expected Final Disposition", "QRE Rules Exercised",
         "Acceptance Scenario", "Why This Path Was Selected",
         "What This Path Adds That Others Do Not", "Tests Hosted",
         "Design State", "Runnable State"],
        [9, 44, 20, 34, 24, 44, 54, 40, 26, 30, 22, 26, 15, 46, 46, 34, 13, 14])

    hosted: dict[str, list[str]] = {}
    for t in targets:
        lt = log_by.get(t.target_id)
        if lt is None:
            continue
        aq = _anchor_question(spec, t)
        ending = None
        if t.dimension == "D2":
            ending = t.subject.partition("<-")[0]
        hp = host_path(paths, aq, ending, _mechanism(spec, t))
        if hp and t.target_id in display_id:
            hosted.setdefault(hp.path_id, []).append(display_id[t.target_id])

    for p in paths:
        # Natural sort, so TC-FT9 comes before TC-FT10 rather than after it.
        ids = sorted(hosted.get(p.path_id, []),
                     key=lambda x: (x[:6], int("".join(ch for ch in x[6:]
                                                      if ch.isdigit()) or 0)))
        _row(ws, [
            p.path_id, p.name, p.route_class, p.preconditions,
            ", ".join(sorted(p.decisions.keys())) or "none",
            "; ".join(f"{k} = {_labels(spec, k, v)}"
                      for k, v in sorted(p.decisions.items())) or "defaults throughout",
            " -> ".join(p.sequence) + (f"  ->  {p.disposition}" if p.disposition else ""),
            ", ".join(p.sequence),
            ", ".join(p.skipped) or "none",
            ", ".join(p.hidden) or "none",
            p.disposition or "-",
            ", ".join(p.rules_exercised) or "none",
            p.scenario or "-",
            p.why_selected,
            p.what_it_adds,
            f"{len(ids)} test(s)" + (": " + ", ".join(ids[:6]) +
                                     (" ..." if len(ids) > 6 else "") if ids else ""),
            "DESIGNED",
            "RUNNABLE" if not p.campaign else "CAMPAIGN",
        ], fonts={1: MONO, 7: MONO, 17: OK})

    # =================== TAB 2: TEST CASES =================================
    ws = _sheet(
        wb, False, "Test Cases",
        "One row per test, ordered the way the questionnaire is ordered: the "
        "screening questions first, then Q1 onward, then anything that spans "
        "the whole survey. Every row names the question it belongs to and the "
        "path it is run on.",
        [c[0] for c in COLUMNS], [c[1] for c in COLUMNS])

    order = 0
    for _, _, _, t, lt, sc, aq in ordered:
        order += 1
        ex = exe_by.get(t.target_id)
        cls = _test_class(t)
        tcid = display_id[t.target_id]

        ending = t.subject.partition("<-")[0] if t.dimension == "D2" else None
        hp = host_path(paths, aq, ending, _mechanism(spec, t))
        obs = sc.expected or {}
        camp = lt.get("campaign") or {}

        setup = "; ".join(
            f"{st['question_id']} = {_labels(spec, st['question_id'], st.get('value'))}"
            for st in (lt.get("setup") or [])) or "none"
        action = lt.get("action_text", "")
        if camp.get("repetitions") and camp.get("cell"):
            pre = (f"{camp['cell']} already at its target of "
                   f"{camp['repetitions']} respondents. Send that run-up first, "
                   f"then run this respondent as number {camp['repetitions'] + 1}.")
        elif camp.get("repetitions"):
            pre = (f"Run this same path {camp['repetitions']} times and compare "
                   f"the recorded option order between runs.")
        elif hp:
            pre = f"Travel path {hp.path_id} ({hp.name}) as far as {aq or 'the end'}."
        else:
            pre = "Fresh respondent, no prior answers."

        seq_text = " -> ".join(obs.get("questions_seen_in_order") or []) or "-"
        if obs.get("blocked_at"):
            seq_text += f"  (refused at {obs['blocked_at']}; respondent stays put)"

        _row(ws, [
            order,
            aq or "whole survey",
            tcid,
            cls,
            PRIMARY_TYPE.get(t.dimension, "-"),
            SECONDARY_TYPE.get((t.dimension, t.polarity), "-"),
            hp.path_id if hp else "not route-based",
            _name(spec, t),
            _objective(spec, t),
            WHY_EXISTS.get(t.dimension, ""),
            PRIORITY.get(t.dimension, "Medium"),
            pre,
            "FRESH_RESPONDENT",
            (setup + ("\nthen: " + action if action else "")),
            TEST_DATA_CLASS.get((t.dimension, t.polarity),
                                "VALID_TYPICAL for every answer"),
            (f"Every answer before {aq} is the first valid option, chosen only "
             f"to reach the question under test without triggering anything "
             f"else. The answer at {aq} is the one the claim turns on."
             if aq else "Chosen to walk the whole survey without triggering a "
                        "screen-out."),
            action or "Answer through the whole survey as listed, then observe.",
            seq_text,
            ", ".join(obs.get("questions_seen_in_order") or []) or "-",
            ", ".join(hp.skipped) if hp and hp.skipped else "none",
            ", ".join(obs.get("questions_not_seen") or []) or "none",
            ("No forward navigation - the same question must reappear."
             if obs.get("blocked_at") else
             "Each listed question is followed immediately by the next in the "
             "sequence."),
            _checks(ex["assertions"] if ex else lt["assertions"]),
            obs.get("ending_reached") or (
                "Not reached - the respondent cannot progress"
                if obs.get("blocked_at") else "-"),
            "Page render per question, plus the ordered sequence log and any "
            "validation message shown.",
            _traceability(spec, t),
            _provenance(spec, t),
            "RESOLVED" if not lt.get("provisional_semantics") else "PROVISIONAL",
            "YES" if ex else "NOT YET",
            _not_proven(spec, t),
            ("Leans on an unconfirmed reading: "
             + ", ".join(lt["provisional_semantics"])
             if lt.get("provisional_semantics") else "-"),
        ], fonts={1: MONO, 2: BOLD, 3: MONO, 7: MONO,
                  28: (OK if not lt.get("provisional_semantics") else WARN),
                  29: (OK if ex else WARN)})

    # =================== TAB 3: NOT ENUMERATED =============================
    ws = _sheet(
        wb, False, "Paths Not Enumerated",
        "Which combinations were deliberately left out of the path set, how "
        "large each family is, and what risk remains. Stating this is what "
        "makes the selected set defensible rather than arbitrary.",
        ["Family", "Size of the family", "Why it was collapsed",
         "Residual risk"],
        [50, 44, 62, 62])
    for item in path_report.get("not_enumerated", []):
        _row(ws, [item["family"], item["combinations"],
                  item["why_collapsed"], item["residual_risk"]])

    _row(ws, ["", "", "", ""])
    _row(ws, ["Branch coverage proof", "", "", ""], fonts={1: BOLD})
    _row(ws, ["Branch decision points in this survey",
              str(path_report["branches_total"]), "", ""])
    _row(ws, ["Branch states needing coverage (each one, both directions)",
              str(path_report["branch_states_total"]), "", ""])
    _row(ws, ["Branch states reached by the selected paths",
              str(path_report["branch_states_covered"]), "", ""],
         fonts={2: OK})
    _row(ws, ["Branch states still unreached",
              str(len(path_report["branch_states_missing"])) or "0",
              ", ".join(path_report["branch_states_missing"]) or "none", ""],
         fonts={2: OK if not path_report["branch_states_missing"] else WARN})

    if path_report.get("branch_states_shadowed"):
        _row(ws, ["", "", "", ""])
        _row(ws, ["Shadowed branches - a finding, not a gap", "", "", ""],
             fonts={1: BOLD})
        for sh in path_report["branch_states_shadowed"]:
            _row(ws, [f"{sh['question']} display rule",
                      f"shadowed by {sh['shadowed_by']}",
                      sh["finding"],
                      "None for testing. Worth raising with the questionnaire "
                      "author, since two rules say the same thing."],
                 fonts={2: WARN})

    wb.save(out_path)
    return out_path
